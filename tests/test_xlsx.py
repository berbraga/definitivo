from openpyxl import load_workbook

from renov_market_scan.report.xlsx import write_xlsx_report

SUMMARY = [
    {
        "erp_code": "20022A0", "device_name": "IPHONE 13 128GB A0", "fabricante": "APPLE",
        "modelo": "IPHONE 13", "capacidade": "128GB", "valor_atual_planilha": 2000.0,
        "n_amostras": 5, "n_fontes": 2, "preco_minimo": 2500.0,
        "link_minimo": "https://olx.com.br/min", "mediana": 2800.0, "preco_maximo": 3200.0,
        "link_maximo": "https://olx.com.br/max", "p25": 2600.0, "p75": 3000.0,
        "spread_pct": 0.25, "razao_mediana_vs_atual": 1.4,
        "condicao_predominante": "seminovo", "fontes": "olx, enjoei",
        "coletado_em": "2026-07-28", "status": "ok", "observacoes": "",
        "min_bruto": 90.0, "max_bruto": 3200.0,
    }
]
SAMPLES = [
    {
        "erp_code": "20022A0", "fonte": "olx", "titulo": "iPhone 13 128GB seminovo",
        "preco": 2800.0, "condicao": "seminovo", "url": "https://olx.com.br/a-1",
        "capturado_em": "2026-07-28T10:00:00", "flag_5g_divergente": "nao",
        "cited_text": "R$ 2.800,00",
    }
]
DISCARDED = [
    {
        "erp_code": "20022A0", "fonte": "olx", "titulo": "Capa iPhone 13", "preco": 30.0,
        "url": "https://olx.com.br/c-1", "motivo_descarte": "acessorio_ou_peca",
        "cited_text": "",
    }
]
ANOMALIES = [
    {
        "linha_planilha": 5, "erp_code": "10000A0", "device_name": "MOTO XT882 1GB A0",
        "campo": "Storage, GB*", "valor_bruto": "1", "motivo": "storage_suspeito",
        "status": "revisao_humana",
    }
]


def write(tmp_path):
    path = tmp_path / "referencia.xlsx"
    write_xlsx_report(path, SUMMARY, SAMPLES, DISCARDED, ANOMALIES, "2026-07-28")
    return path


def test_the_four_sheets_exist_with_the_expected_names(tmp_path):
    workbook = load_workbook(write(tmp_path))
    assert workbook.sheetnames == ["Resumo", "Amostras", "Descartados", "Anomalias"]
    workbook.close()


def test_headers_are_written_and_data_starts_on_row_two(tmp_path):
    workbook = load_workbook(write(tmp_path))
    sheet = workbook["Resumo"]
    assert sheet.cell(row=1, column=1).value == "erp_code"
    assert sheet.cell(row=2, column=1).value == "20022A0"
    workbook.close()


def test_links_are_written_as_clickable_formulas(tmp_path):
    workbook = load_workbook(write(tmp_path))
    sheet = workbook["Resumo"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("link_minimo") + 1
    value = sheet.cell(row=2, column=column).value
    assert isinstance(value, str) and value.startswith("=HYPERLINK(")
    assert "https://olx.com.br/min" in value
    workbook.close()


def test_freeze_pane_and_autofilter_are_set(tmp_path):
    workbook = load_workbook(write(tmp_path))
    sheet = workbook["Resumo"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    workbook.close()


def test_currency_columns_carry_the_brl_number_format(tmp_path):
    workbook = load_workbook(write(tmp_path))
    sheet = workbook["Resumo"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("mediana") + 1
    assert "R$" in sheet.cell(row=2, column=column).number_format
    workbook.close()


def test_the_footnote_is_written_below_the_summary(tmp_path):
    workbook = load_workbook(write(tmp_path))
    sheet = workbook["Resumo"]
    texts = [
        sheet.cell(row=row, column=1).value
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    ]
    assert any("nao de transacao" in str(text) for text in texts)
    workbook.close()


def test_empty_sheets_still_get_their_headers(tmp_path):
    path = tmp_path / "vazio.xlsx"
    write_xlsx_report(path, SUMMARY, [], [], [], "2026-07-28")
    workbook = load_workbook(path)
    assert workbook["Amostras"].cell(row=1, column=1).value == "erp_code"
    assert workbook["Amostras"].max_row == 1
    workbook.close()
