import hashlib

from openpyxl import load_workbook

from renov_market_scan.ingest.reader import EXPECTED_HEADER
from renov_market_scan.report.template_copy import EXTRA_COLUMNS, write_template_copy

SUMMARY_BY_KEY = {
    ("10870A0", "GALAXY A17", "128GB"): {
        "n_amostras": 5, "preco_minimo": 300.0, "mediana": 400.0, "preco_maximo": 500.0,
        "link_minimo": "https://olx.com.br/min", "link_maximo": "https://olx.com.br/max",
        "razao_mediana_vs_atual": 1.0, "status": "ok",
    }
}


def test_the_input_file_is_untouched(tmp_path, make_sheet, device_row_dict):
    source = make_sheet(tmp_path, [device_row_dict()])
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_the_copy_preserves_the_help_row_header_and_sheet_name(
    tmp_path, make_sheet, device_row_dict
):
    source = make_sheet(tmp_path, [device_row_dict()], sheet_name="iPhones")
    destination = write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    workbook = load_workbook(destination)
    assert workbook.sheetnames[0] == "iPhones"
    sheet = workbook["iPhones"]
    assert "Custom name of a device" in str(sheet.cell(row=1, column=1).value)
    header = [sheet.cell(row=2, column=index + 1).value for index in range(len(EXPECTED_HEADER))]
    assert tuple(header) == EXPECTED_HEADER
    workbook.close()


def test_extra_columns_are_appended_to_the_right(tmp_path, make_sheet, device_row_dict):
    source = make_sheet(tmp_path, [device_row_dict()])
    destination = write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    workbook = load_workbook(destination)
    sheet = workbook[workbook.sheetnames[0]]
    first_extra = len(EXPECTED_HEADER) + 1
    appended = [
        sheet.cell(row=2, column=first_extra + offset).value
        for offset in range(len(EXTRA_COLUMNS))
    ]
    assert tuple(appended) == EXTRA_COLUMNS
    workbook.close()


def test_matching_rows_receive_their_values(tmp_path, make_sheet, device_row_dict):
    source = make_sheet(tmp_path, [device_row_dict()])
    destination = write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    workbook = load_workbook(destination)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = 2
    header = [cell.value for cell in sheet[header_row]]
    column = header.index("mediana_mercado") + 1
    assert sheet.cell(row=3, column=column).value == 400.0
    workbook.close()


def test_the_original_price_column_is_never_overwritten(tmp_path, make_sheet, device_row_dict):
    source = make_sheet(tmp_path, [device_row_dict()])
    destination = write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    workbook = load_workbook(destination)
    sheet = workbook[workbook.sheetnames[0]]
    column = EXPECTED_HEADER.index("Price for In-store") + 1
    assert sheet.cell(row=3, column=column).value == 400
    workbook.close()


def test_rows_without_a_result_are_left_blank(tmp_path, make_sheet, device_row_dict):
    source = make_sheet(tmp_path, [device_row_dict(**{"ERP Code": "SEM_RESULTADO"})])
    destination = write_template_copy(source, tmp_path / "copia.xlsx", SUMMARY_BY_KEY)
    workbook = load_workbook(destination)
    sheet = workbook[workbook.sheetnames[0]]
    header = [cell.value for cell in sheet[2]]
    column = header.index("mediana_mercado") + 1
    assert sheet.cell(row=3, column=column).value is None
    workbook.close()
