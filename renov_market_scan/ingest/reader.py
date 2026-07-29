"""Read the device import template into DeviceRow objects.

The input file is opened read-only and is never written to.
"""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from renov_market_scan.models import DeviceRow

EXPECTED_HEADER: tuple[str, ...] = (
    "Device name*",
    "Manufacturer*",
    "Model*",
    "Device type*",
    "Storage, GB*",
    "Ram, GB*",
    "Color*",
    "Price for In-store",
    "Price for Widget & Mobile",
    "Buying",
    "Questionnaire Widget",
    "Questionnaire In-store",
    "Questionnaire Mobile",
    "Questionnaire robot",
    "Priority",
    "Reward Types",
    "Minimum Price",
    "Maximum Price",
    "ERP Code",
)

HEADER_ROW = 2
FIRST_DATA_ROW = 3


class SheetFormatError(Exception):
    """The workbook does not match the expected import template."""


def _text(value: Any) -> str:
    """Normalize a cell to a trimmed string. Integers lose the float tail."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any) -> float | None:
    """Return the cell as a float, or None when it is not numeric."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def find_data_sheet(workbook: Workbook) -> str:
    """Return the sheet whose row 2 contains 'Device name*'.

    Falls back to the first sheet. The sheet name is never hardcoded.
    """
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        rows = list(worksheet.iter_rows(min_row=1, max_row=HEADER_ROW, values_only=True))
        if len(rows) >= HEADER_ROW and any(
            _text(cell) == "Device name*" for cell in rows[HEADER_ROW - 1]
        ):
            return name
    return workbook.sheetnames[0]


def read_device_rows(path: Path) -> tuple[list[DeviceRow], list[str]]:
    """Parse the template. Returns (rows, warnings) with warnings in pt-BR."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = find_data_sheet(workbook)
        worksheet = workbook[sheet_name]
        raw_rows = list(worksheet.iter_rows(values_only=True))

        if len(raw_rows) < HEADER_ROW:
            raise SheetFormatError(
                f"A aba {sheet_name!r} tem menos de {HEADER_ROW} linhas; "
                "o template exige linha de ajuda e cabecalho."
            )

        header = tuple(_text(cell) for cell in raw_rows[HEADER_ROW - 1] if _text(cell))
        if header != EXPECTED_HEADER:
            missing = [c for c in EXPECTED_HEADER if c not in header]
            extra = [c for c in header if c not in EXPECTED_HEADER]
            raise SheetFormatError(
                f"Cabecalho da aba {sheet_name!r} nao corresponde ao template. "
                f"Faltando: {missing}. Inesperadas: {extra}."
            )

        index = {name: position for position, name in enumerate(header)}
        rows: list[DeviceRow] = []
        warnings: list[str] = []

        for offset, raw in enumerate(raw_rows[HEADER_ROW:], start=FIRST_DATA_ROW):
            if not any(_text(cell) for cell in raw):
                continue

            def cell(column: str, source: tuple[Any, ...] = raw) -> Any:
                position = index[column]
                return source[position] if position < len(source) else None

            instore = _number(cell("Price for In-store"))
            widget = _number(cell("Price for Widget & Mobile"))
            if instore != widget:
                warnings.append(
                    f"linha {offset}: 'Price for In-store' ({instore}) difere de "
                    f"'Price for Widget & Mobile' ({widget}); usando In-store."
                )

            storage = _text(cell("Storage, GB*"))
            ram = _text(cell("Ram, GB*"))
            rows.append(
                DeviceRow(
                    row_number=offset,
                    device_name=_text(cell("Device name*")),
                    manufacturer=_text(cell("Manufacturer*")),
                    model=_text(cell("Model*")),
                    device_type=_text(cell("Device type*")),
                    storage_raw=storage or None,
                    ram_raw=ram or None,
                    color=_text(cell("Color*")),
                    price_instore=instore,
                    price_widget_mobile=widget,
                    erp_code=_text(cell("ERP Code")),
                )
            )
        return rows, warnings
    finally:
        workbook.close()
