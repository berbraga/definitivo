"""Write the four-sheet auditable report."""

from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScale, FormatObject, Rule
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from renov_market_scan.report.assemble import (
    ANOMALY_COLUMNS,
    DISCARDED_COLUMNS,
    FOOTNOTE,
    SAMPLE_COLUMNS,
    SUMMARY_COLUMNS,
)

BRL_FORMAT = "R$ #,##0.00"
RATIO_FORMAT = "0.00"

CURRENCY_COLUMNS: frozenset[str] = frozenset(
    {
        "valor_atual_planilha",
        "preco_minimo",
        "mediana",
        "preco_maximo",
        "p25",
        "p75",
        "min_bruto",
        "max_bruto",
        "preco",
    }
)
LINK_COLUMNS: frozenset[str] = frozenset({"link_minimo", "link_maximo", "url"})
RATIO_COLUMNS: frozenset[str] = frozenset({"razao_mediana_vs_atual", "spread_pct"})

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10


def _hyperlink(url: str) -> str:
    """A clickable cell. Quotes in the URL are escaped for the formula."""
    safe = url.replace('"', '""')
    return f'=HYPERLINK("{safe}";"abrir")'


def _write_sheet(sheet: Worksheet, columns: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    sheet.append(list(columns))
    for row in rows:
        sheet.append(
            [
                _hyperlink(str(row[column]))
                if column in LINK_COLUMNS and row.get(column)
                else row.get(column)
                for column in columns
            ]
        )

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        longest = max([len(str(column))] + [len(str(row.get(column) or "")) for row in rows])
        sheet.column_dimensions[letter].width = min(
            max(longest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH
        )
        if column in CURRENCY_COLUMNS:
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row=row_index, column=index).number_format = BRL_FORMAT
        elif column in RATIO_COLUMNS:
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row=row_index, column=index).number_format = RATIO_FORMAT


def write_xlsx_report(
    path: Path,
    summary_rows: list[dict[str, Any]],
    sample_rows: list[dict[str, Any]],
    discarded_rows: list[dict[str, Any]],
    anomaly_rows: list[dict[str, Any]],
    collected_on: str,
) -> None:
    """Write Resumo, Amostras, Descartados and Anomalias, in that order."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    summary = cast(Worksheet, workbook.active)
    summary.title = "Resumo"
    _write_sheet(summary, SUMMARY_COLUMNS, summary_rows)

    if "razao_mediana_vs_atual" in SUMMARY_COLUMNS and summary_rows:
        index = SUMMARY_COLUMNS.index("razao_mediana_vs_atual") + 1
        letter = get_column_letter(index)
        color_scale = ColorScale(
            cfvo=[
                FormatObject(type="min"),
                FormatObject(type="num", val=1),
                FormatObject(type="max"),
            ],
            color=[Color("F8696B"), Color("FFEB84"), Color("63BE7B")],
        )
        summary.conditional_formatting.add(
            f"{letter}2:{letter}{len(summary_rows) + 1}",
            Rule(type="colorScale", colorScale=color_scale),
        )

    footnote_row = len(summary_rows) + 3
    summary.cell(row=footnote_row, column=1).value = FOOTNOTE.format(data=collected_on)

    _write_sheet(workbook.create_sheet("Amostras"), SAMPLE_COLUMNS, sample_rows)
    _write_sheet(workbook.create_sheet("Descartados"), DISCARDED_COLUMNS, discarded_rows)
    _write_sheet(workbook.create_sheet("Anomalias"), ANOMALY_COLUMNS, anomaly_rows)

    workbook.save(path)
    workbook.close()
